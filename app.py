import re, io
from pathlib import Path
import pandas as pd
import streamlit as st
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.pipeline import Pipeline
from sklearn.linear_model import LogisticRegression

DATA = Path(__file__).parent / "data" / "ContractIQ_clause_dataset.csv"
st.set_page_config(page_title="ContractIQ", page_icon="📄", layout="wide")
st.title("📄 ContractIQ")
st.subheader("AI-assisted contract clause classification and risk review")
st.caption("Educational prototype only. Predictions are not legal advice.")

def clean(s):
    return re.sub(r"\\s+", " ", str(s or "").lower()).strip()

def extract(f):
    name, raw = f.name.lower(), f.getvalue()
    if name.endswith(".txt"):
        return raw.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(raw)).pages)
    if name.endswith(".docx"):
        from docx import Document
        d = Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        return "\\n".join(parts)
    raise ValueError("Upload a PDF, DOCX, or TXT file.")

def split_clauses(text):
    parts = [re.sub(r"\\s+", " ", p).strip() for p in re.split(r"\\n\\s*\\n", text) if p.strip()]
    if len(parts) <= 1:
        parts = [p.strip() for p in re.split(r"(?<=[.;])\\s+(?=[A-Z])", text) if p.strip()]
    return [p for p in parts if len(p) >= 20]

@st.cache_resource
def train():
    if not DATA.exists():
        raise FileNotFoundError("Place ContractIQ_clause_dataset.csv in data/ folder.")
    df = pd.read_csv(DATA)
    required = {"clause_text", "clause_type", "risk_level"}
    if not required.issubset(df.columns):
        raise ValueError("Dataset must contain: clause_text, clause_type, risk_level")
    df = df.dropna(subset=list(required)).copy()
    df["clause_text"] = df["clause_text"].map(clean)
    models = {}
    for target in ["clause_type", "risk_level"]:
        model = Pipeline([
            ("tfidf", TfidfVectorizer(max_features=3000, ngram_range=(1, 2), sublinear_tf=True)),
            ("classifier", LogisticRegression(max_iter=1500, class_weight="balanced"))
        ])
        model.fit(df["clause_text"], df[target].astype(str))
        models[target] = model
    return models, len(df)

try:
    models, rows = train()
    st.sidebar.success(f"Training rows loaded: {rows}")
except Exception as e:
    st.error(f"Model setup failed: {e}")
    st.stop()

tab1, tab2, tab3 = st.tabs(["Analyze contract", "Test a clause", "About"])
with tab1:
    st.write("Upload a contract. Text-based PDFs work best; scanned PDFs require OCR.")
    f = st.file_uploader("Contract file", type=["pdf", "docx", "txt"])
    if f:
        try:
            text = extract(f)
            with st.expander("Extracted text preview"):
                st.text(text[:8000])
            clauses = split_clauses(text)
            if not clauses:
                clauses = [x.strip() for x in text.splitlines() if len(x.strip()) >= 20]
            if not clauses:
                st.warning("No clause text detected. Try a text-based document.")
            else:
                results = []
                for i, clause in enumerate(clauses, 1):
                    row = {"Clause #": i, "Clause text": clause}
                    for target, label in [("clause_type", "Predicted clause type"), ("risk_level", "Predicted risk")]:
                        model = models[target]
                        row[label] = model.predict([clean(clause)])[0]
                        row[label + " confidence"] = round(float(max(model.predict_proba([clean(clause)])[0])), 3)
                    results.append(row)
                out = pd.DataFrame(results)
                a, b, c = st.columns(3)
                a.metric("Clauses detected", len(out))
                b.metric("Predicted high-risk clauses", int(out["Predicted risk"].str.lower().eq("high").sum()))
                c.metric("Clause categories", out["Predicted clause type"].nunique())
                st.dataframe(out, use_container_width=True, hide_index=True)
                st.download_button("Download results as CSV", out.to_csv(index=False).encode(), "contractiq_analysis.csv", "text/csv")
                st.warning("Review every result with a qualified person. Model confidence is not legal certainty.")
        except Exception as e:
            st.error(f"Could not analyze the file: {e}")

with tab2:
    clause = st.text_area("Paste a clause", height=160, placeholder="The customer shall pay all undisputed invoices within 30 days.")
    if st.button("Analyze clause", type="primary"):
        if len(clause.strip()) < 10:
            st.warning("Enter a longer clause.")
        else:
            x = clean(clause)
            a, b = st.columns(2)
            a.metric("Clause type", models["clause_type"].predict([x])[0])
            b.metric("Risk level", models["risk_level"].predict([x])[0])
            with st.expander("Confidence details"):
                for target in ["clause_type", "risk_level"]:
                    probs = models[target].predict_proba([x])[0]
                    st.write(target.replace("_", " ").title(), f"{max(probs):.1%}")

with tab3:
    st.markdown("""
    **Workflow:** CSV dataset → TF-IDF text features → Logistic Regression → contract text extraction → clause segmentation → clause type and risk predictions → CSV export.

    **Limitations:** this starter version uses heuristic clause splitting, does not perform OCR on scanned PDFs, and does not include RAG, an external LLM, a vector database, or autonomous agents. Use approved systems for confidential contracts.
    """)
st.divider()
st.caption("ContractIQ • Local prototype")
